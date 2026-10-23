#!/usr/bin/env python3
"""Build the BADM 534 resource library catalogs from the YAML source files.

Usage:  python3 build_catalog.py <01_Resources folder>

Reads   00_Index/_data/*.yaml          (source of truth; edit these, then rebuild)
Writes  <folder>/README.md             (one annotated catalog per category folder)
        00_Index/README.md             (master index: overview, core picks, module x AI crosswalk)
        00_Index/resource_catalog.csv  (flat table)
        00_Index/resource_catalog.xlsx (filterable spreadsheet)
        00_Index/resource_catalog.json (structured data for the future wiki)
"""
import csv
import json
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date
from pathlib import Path

import yaml

ROOT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2]
DATA = ROOT / "00_Index" / "_data"
if __name__ == "__main__" and not DATA.is_dir():
    sys.exit(f"Expected the 01_Resources folder (containing 00_Index/_data); got {ROOT}")

FOLDERS = OrderedDict([
    ("M1", ("01_M1_Marketing_Analytics_and_Python", "M1 · Introduction to Marketing Analytics and Python")),
    ("M2", ("02_M2_Customer_Data_and_Visualization", "M2 · Customer Data and Visualization")),
    ("M3", ("03_M3_Regression", "M3 · Regression Analysis: A Review")),
    ("M4", ("04_M4_Conjoint_Segmentation_CLV", "M4 · Conjoint Analysis, Clustering and Customer Lifetime Value")),
    ("M5", ("05_M5_Text_and_Social_Media_Analytics", "M5 · Content Analysis, Text and Social Media")),
    ("M6", ("06_M6_Causal_Inference_and_Experiments", "M6 · Causal Models and Experiments")),
    ("M7", ("07_M7_Predictive_Models_ML_and_AI", "M7 · Predictive Models, AI and Machine Learning")),
    ("M8", ("08_M8_Current_Topics_in_Marketing", "M8 · Current Topics in Marketing")),
    ("AI_EVIDENCE", ("10_AI_Coding_Concepts_and_Evidence", "AI-Assisted Coding · Concepts, History and Evidence")),
    ("AI_TOOLS", ("11_AI_Coding_Tools_and_Workflows", "AI-Assisted Coding · Tools and Workflows")),
    ("AI_PEDAGOGY", ("12_AI_Coding_Pedagogy_and_Policy", "AI-Assisted Coding · Teaching, Learning and Policy")),
    ("DATASETS", ("20_Datasets", "Datasets for Exercises and Projects")),
    ("TEXTBOOKS", ("21_Open_Textbooks_and_Courses", "Open Textbooks and Courses (cross-module)")),
    ("WIKI", ("22_Wiki_Design_and_Hosting", "Interactive Wiki · Design and Hosting")),
])
MODULE_TAG_TO_CAT = {f"M{i}": f"M{i}" for i in range(1, 9)}
MODULE_TAG_TO_CAT["WIKI"] = "WIKI"
PRIORITY_ORDER = {"core": 0, "recommended": 1, "optional": 2}
FIELDS = ["id", "folder", "category", "section", "subtopic", "priority", "title", "authors", "year", "type",
          "venue", "url", "doi", "oa_url", "access", "license", "level", "modules", "ai_angle",
          "summary", "wiki_use", "tools", "local_copy", "verified_via", "notes"]


def norm_url(u):
    u = (u or "").strip().lower()
    u = re.sub(r"^https?://(www\.)?", "", u)
    return u.rstrip("/")


def as_list(x):
    if x is None:
        return []
    if isinstance(x, (list, tuple)):
        return [str(i) for i in x]
    return [s.strip() for s in str(x).split(",") if s.strip()]


def load_sections():
    """_data/sections.yaml maps each category to ordered sections, each listing the raw subtopic labels it absorbs."""
    f = DATA / "sections.yaml"
    if not f.exists():
        return {}
    raw = yaml.safe_load(f.read_text()) or {}
    out = {}
    for cat, secs in raw.items():
        out[cat] = OrderedDict()
        for sec in secs or []:
            for name, subs in sec.items():
                out[cat][name] = [str(x).strip().lower() for x in (subs or [])]
    return out


SECTIONS = {}


def section_of(e):
    if e.get("section"):
        return str(e["section"])
    sub = str(e.get("subtopic") or "General").strip()
    for name, subs in SECTIONS.get(e["category"], {}).items():
        if sub.lower() in subs:
            return name
    return sub


KEEP_ORDER = ["DATASETS", "TEXTBOOKS", "AI_EVIDENCE", "AI_TOOLS", "AI_PEDAGOGY",
              "M1", "M2", "M3", "M4", "M5", "M6", "M7", "M8", "WIKI"]


def norm_title(t):
    t = re.sub(r"\([^)]*\)", "", str(t).lower())          # drop parentheticals like "(X post)"
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def load_curation():
    f = DATA / "curation.yaml"
    c = yaml.safe_load(f.read_text()) if f.exists() else {}
    c = c or {}
    return {"prefer": set(c.get("prefer") or []), "drop": set(c.get("drop") or []),
            "same": [list(x) for x in (c.get("same") or [])], "distinct": [set(x) for x in (c.get("distinct") or [])],
            "set": c.get("set") or {}}


def load():
    cur = load_curation()
    raw, seen_id = [], set()
    for f in sorted(DATA.glob("*.yaml")):
        if f.name in ("sections.yaml", "curation.yaml"):
            continue
        for e in yaml.safe_load(f.read_text()) or []:
            if not isinstance(e, dict) or not e.get("title") or not e.get("url") or e.get("category") not in FOLDERS:
                continue
            e["modules"] = as_list(e.get("modules"))
            e["tools"] = as_list(e.get("tools"))
            base, n = e.get("id") or re.sub(r"\W+", "-", e["title"].lower())[:40], 2
            eid = base
            while eid in seen_id:
                eid, n = f"{base}-{n}", n + 1
            e["id"] = eid
            seen_id.add(eid)
            e.update(cur["set"].get(eid, {}))           # manual field corrections
            if eid in cur["drop"]:
                continue
            raw.append(e)
    # union-find over duplicate keys: same URL, same normalised title, or declared "same"
    parent = {e["id"]: e["id"] for e in raw}

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b):
        if {a, b} in cur["distinct"]:
            return
        parent[find(a)] = find(b)

    by_key = {}
    for e in raw:
        for k in ("u:" + norm_url(e["url"]), "t:" + norm_title(e["title"])):
            if k in by_key:
                union(e["id"], by_key[k])
            else:
                by_key[k] = e["id"]
    for grp in cur["same"]:
        ids = [i for i in grp if i in parent]
        for a in ids[1:]:
            union(a, ids[0])
    groups = OrderedDict()
    for e in raw:
        groups.setdefault(find(e["id"]), []).append(e)
    entries, dupes = [], []
    for grp in groups.values():
        grp.sort(key=lambda e: (e["id"] not in cur["prefer"], KEEP_ORDER.index(e["category"])))
        keep = grp[0]
        for d in grp[1:]:
            mods = set(keep["modules"]) | set(d["modules"])
            if d["category"].startswith("M") or d["category"] == "WIKI":
                mods.add(d["category"])
            keep["modules"] = sorted(mods)
            for fld in ("doi", "oa_url", "license", "download", "venue", "authors", "year"):
                if not keep.get(fld) and d.get(fld):
                    keep[fld] = d[fld]
            if not keep.get("oa_url") and keep.get("access") != "open" and d.get("access") == "open" and norm_url(d["url"]) != norm_url(keep["url"]):
                keep["oa_url"] = d["url"]
            if PRIORITY_ORDER.get(d.get("priority"), 3) < PRIORITY_ORDER.get(keep.get("priority"), 3):
                keep["priority"] = d["priority"]
            keep["ai_angle"] = bool(keep.get("ai_angle") or d.get("ai_angle"))
            dupes.append((d["id"], keep["id"]))
        keep["folder"] = FOLDERS[keep["category"]][0]
        keep["section"] = section_of(keep)
        entries.append(keep)
    return entries, dupes


def fmt_entry(e, local_root=None):
    title = e["title"].strip()
    yr = f" ({e['year']})" if e.get("year") else ""
    line = f"- **[{title}]({e['url']})**{yr}"
    by = []
    if e.get("authors"):
        by.append(str(e["authors"]).strip().rstrip("."))
    if e.get("venue"):
        by.append(f"*{e['venue']}*")
    if by:
        line += " — " + ". ".join(by)
    tags = [e.get("type"), e.get("level"), f"access: {e.get('access')}" if e.get("access") else None,
            e.get("license"), "AI angle" if e.get("ai_angle") else None]
    tags = [f"`{t}`" for t in tags if t]
    out = [line, "  " + " · ".join(tags) if tags else None]
    if e.get("summary"):
        out.append(f"  {str(e['summary']).strip()}")
    if e.get("wiki_use"):
        out.append(f"  *Use in wiki:* {str(e['wiki_use']).strip()}")
    extras = []
    if e.get("doi"):
        d = str(e["doi"]).strip()
        extras.append(f"DOI [{d}](https://doi.org/{d})")
    if e.get("oa_url"):
        extras.append(f"[Open-access copy]({e['oa_url']})")
    if e.get("tools"):
        extras.append("Tools: " + ", ".join(e["tools"]))
    if e.get("local_copy"):
        extras.append(f"Local copy: `{e['local_copy']}`")
    if extras:
        out.append("  " + " · ".join(extras))
    if e.get("notes"):
        out.append(f"  *Note:* {str(e['notes']).strip()}")
    return "\n".join(o for o in out if o)


def sort_key(e):
    return (PRIORITY_ORDER.get(str(e.get("priority", "optional")).lower(), 3), -(int(e["year"]) if str(e.get("year", "")).isdigit() else 0), e["title"].lower())


def write_folder_readmes(entries):
    by_cat = defaultdict(list)
    for e in entries:
        by_cat[e["category"]].append(e)
    for cat, (folder, heading) in FOLDERS.items():
        items = by_cat.get(cat, [])
        d = ROOT / folder
        d.mkdir(exist_ok=True)
        n_core = sum(1 for e in items if e.get("priority") == "core")
        lines = [f"# {heading}", "",
                 f"{len(items)} resources ({n_core} core). Built {date.today().isoformat()} from "
                 f"`00_Index/_data/`; edit the YAML there and rerun `00_Index/_build/build_catalog.py` rather than editing this file.",
                 "", "Priority: **core** = likely to assign or build on; **recommended** = strong supplement; "
                 "**optional** = background or niche. *Access* tells you whether students can open it without a login.", ""]
        subs = OrderedDict()
        for e in items:
            subs.setdefault(e["section"], []).append(e)
        declared = list(SECTIONS.get(cat, {}))
        order = [s for s in declared if s in subs] + sorted(
            (s for s in subs if s not in declared),
            key=lambda s: (-sum(1 for e in subs[s] if e.get("priority") == "core"), -len(subs[s]), s.lower()))
        if len(order) > 1:
            lines += ["## Contents", ""] + [f"- [{s}](#{re.sub(r'[^a-z0-9 -]', '', s.lower()).replace(' ', '-')}) ({len(subs[s])})" for s in order] + [""]
        for s in order:
            lines += [f"## {s}", ""]
            for e in sorted(subs[s], key=sort_key):
                pr = str(e.get("priority") or "optional").lower()
                lines.append(fmt_entry(e).replace("- **[", f"- [{pr}] **[", 1))
                lines.append("")
        if cat.startswith("M"):
            ai = [e for e in items if e.get("ai_angle")]
            cross = [e for e in entries if e["category"] != cat and cat in e["modules"]]
            if ai:
                lines += ["## AI-assisted coding angle for this module", "",
                          "Resources in this folder that show the method done with AI help, LLM versions of the method, or documented failure modes of AI-written analysis:", ""]
                lines += [f"- [{e['title']}]({e['url']})" for e in sorted(ai, key=sort_key)] + [""]
            if cross:
                lines += ["## Also relevant (catalogued in other folders)", ""]
                grouped = OrderedDict()
                for e in sorted(cross, key=lambda e: (list(FOLDERS).index(e["category"]), sort_key(e))):
                    grouped.setdefault(e["folder"], []).append(e)
                for fold, es in grouped.items():
                    lines.append(f"- **`{fold}`**: " + "; ".join(f"[{e['title']}]({e['url']})" for e in es))
                lines.append("")
        (d / "README.md").write_text("\n".join(lines).rstrip() + "\n")


def write_tables(entries):
    idx = ROOT / "00_Index"
    rows = []
    for e in sorted(entries, key=lambda e: (list(FOLDERS).index(e["category"]), e["section"], sort_key(e))):
        r = {k: e.get(k, "") for k in FIELDS}
        r["modules"] = ", ".join(e["modules"])
        r["tools"] = ", ".join(e["tools"])
        r["ai_angle"] = "yes" if e.get("ai_angle") else ""
        rows.append({k: ("" if v is None else v) for k, v in r.items()})
    with open(idx / "resource_catalog.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    (idx / "resource_catalog.json").write_text(json.dumps(
        {"built": date.today().isoformat(), "folders": {k: v[0] for k, v in FOLDERS.items()},
         "resources": [{k: e.get(k) for k in FIELDS + ["download"] if e.get(k) not in (None, "", [])} for e in entries]},
        indent=1, default=str))
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Catalog"
    ws.append(FIELDS)
    for r in rows:
        ws.append([r[k] for k in FIELDS])
    widths = {"id": 22, "folder": 30, "category": 13, "section": 26, "subtopic": 24, "priority": 12, "title": 50, "authors": 28,
              "year": 7, "type": 12, "venue": 22, "url": 40, "doi": 20, "oa_url": 30, "access": 14, "license": 14,
              "level": 12, "modules": 14, "ai_angle": 9, "summary": 60, "wiki_use": 50, "tools": 24,
              "local_copy": 30, "verified_via": 10, "notes": 40}
    for i, k in enumerate(FIELDS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 15)
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = PatternFill("solid", fgColor="13294B")
    url_col = FIELDS.index("url") + 1
    for row in range(2, ws.max_row + 1):
        c = ws.cell(row, url_col)
        if c.value:
            c.hyperlink = c.value
            c.font = Font(color="1F4E9E", underline="single")
        for col in (FIELDS.index("summary") + 1, FIELDS.index("wiki_use") + 1, FIELDS.index("title") + 1):
            ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = ws.dimensions
    # summary sheet
    s = wb.create_sheet("Summary")
    s.append(["folder", "resources", "core", "open access", "AI angle"])
    for cat, (folder, _) in FOLDERS.items():
        it = [e for e in entries if e["category"] == cat]
        s.append([folder, len(it), sum(e.get("priority") == "core" for e in it),
                  sum(e.get("access") == "open" for e in it), sum(bool(e.get("ai_angle")) for e in it)])
    s.column_dimensions["A"].width = 42
    for c in s[1]:
        c.font = Font(bold=True)
    wb.save(idx / "resource_catalog.xlsx")


def write_index(entries, dupes):
    idx = ROOT / "00_Index"
    tot = len(entries)
    lines = ["# BADM 534 Resource Library — Master Index", "",
             f"{tot} catalogued resources across {len(FOLDERS)} folders · built {date.today().isoformat()}.", "",
             "The course spine is the previous syllabus's eight modules (M1–M8). On top of that sit a dedicated AI-assisted coding track (folders 10–12), "
             "a datasets folder, cross-module open textbooks, and resources for designing and hosting the interactive wiki.", "",
             "**How this is organised**", "",
             "- Each folder's `README.md` is an annotated catalog grouped by subtopic and ranked core → recommended → optional.",
             "- `resource_catalog.xlsx` / `.csv` hold the same entries in one filterable table. `resource_catalog.json` is the structured version the wiki can be generated from later.",
             "- `_data/*.yaml` is the source of truth. Edit an entry there, then run `python3 _build/build_catalog.py` (from `00_Index`) to regenerate everything.",
             "- `repos/` subfolders inside a module folder hold local copies of openly licensed GitHub material (notebooks, code, small datasets). Each one includes a `SOURCE.md` giving its origin and license.",
             ""]
    intro = DATA / "index_intro.md"
    if intro.exists():
        by_id = {e["id"]: e for e in entries}
        alias = dict(dupes)

        def link(m):
            i = m.group(1)
            e = by_id.get(i) or by_id.get(alias.get(i, ""))
            if not e:
                print(f"warning: index_intro references unknown id {i}")
                return i
            return f"[{e['title']}]({e['url']})"
        lines += [re.sub(r"\{\{([\w.-]+)\}\}", link, intro.read_text()).rstrip(), ""]
    lines += ["## Folders", "", "| Folder | Resources | Core | Open access | AI angle |", "|---|---:|---:|---:|---:|"]
    for cat, (folder, heading) in FOLDERS.items():
        it = [e for e in entries if e["category"] == cat]
        lines.append(f"| [{heading}](../{folder}/README.md) | {len(it)} | {sum(e.get('priority') == 'core' for e in it)} | "
                     f"{sum(e.get('access') == 'open' for e in it)} | {sum(bool(e.get('ai_angle')) for e in it)} |")
    lines += ["", "## Module × AI-assisted coding crosswalk", "",
              "For each module, the core resources together with the resources that carry an AI-assisted coding angle. Use this as a starting point when designing each wiki page.", ""]
    for i in range(1, 9):
        m = f"M{i}"
        tagged = [e for e in entries if e["category"] == m or m in e["modules"]]
        core = [e for e in tagged if e.get("priority") == "core"]
        ai = [e for e in tagged if e.get("ai_angle")]
        lines += [f"### {FOLDERS[m][1]}", ""]
        lines += ["*Core:* " + "; ".join(f"[{e['title']}]({e['url']})" for e in sorted(core, key=sort_key)[:12]), ""]
        if ai:
            lines += ["*AI angle:* " + "; ".join(f"[{e['title']}]({e['url']})" for e in sorted(ai, key=sort_key)[:10]), ""]
    if dupes:
        lines += ["## Merged duplicates", "", "These entries turned up in more than one search. Each was kept once, and its module tags were merged:", ""]
        lines += [f"- `{a}` → `{b}`" for a, b in dupes] + [""]
    (idx / "README.md").write_text("\n".join(lines).rstrip() + "\n")


if __name__ == "__main__":
    SECTIONS.update(load_sections())
    entries, dupes = load()
    # local copies recorded by the download step
    lc = ROOT / "00_Index" / "_data" / "local_copies.json"
    if lc.exists():
        m = json.loads(lc.read_text())
        for dropped, kept in dupes:                 # a merged duplicate's local copy belongs to the kept entry
            if dropped in m and kept not in m:
                m[kept] = m[dropped]
        for e in entries:
            if e["id"] in m:
                e["local_copy"] = m[e["id"]]
    write_folder_readmes(entries)
    write_tables(entries)
    write_index(entries, dupes)
    print(f"{len(entries)} entries, {len(dupes)} duplicates merged")
